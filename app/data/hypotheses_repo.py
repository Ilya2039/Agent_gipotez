from __future__ import annotations

import json
import logging
import os
from pathlib import Path
from typing import Dict, List, Tuple, Iterable
import html as _html

from docx import Document  # type: ignore


DATA_DIR = Path("data")
HYPOTHESES_TXT = DATA_DIR / "hypotheses.txt"
HYPOTHESES_DOCX = DATA_DIR / "SP_Порядок_проработки_развилок_и_гипотез_2025_10_09.docx"
HYPOTHESES_MAP_JSON = DATA_DIR / "hypotheses_map.json"
HYPOTHESES_FULL_JSON = DATA_DIR / "hypotheses_full.json"
HYPOTHESES_CARDS_JSON = DATA_DIR / "hypotheses_cards.json"


def load_hypotheses_list() -> List[str]:
    lines = (HYPOTHESES_TXT.read_text(encoding="utf-8") if HYPOTHESES_TXT.exists() else "").splitlines()
    return [ln.strip() for ln in lines if ln.strip()]


def _normalize_title(s: str) -> str:
    import re

    s = s.strip().lower()
    s = re.sub(r"\s+", " ", s)
    s = s.replace("—", "-")
    return s


def _best_title_match(norm_text: str, title_norm_to_orig: Dict[str, str]) -> Tuple[str | None, float]:
    # exact
    if norm_text in title_norm_to_orig:
        return title_norm_to_orig[norm_text], 1.0
    # fuzzy
    try:
        from difflib import SequenceMatcher
    except Exception:
        return None, 0.0
    best: Tuple[str | None, float] = (None, 0.0)
    for norm_title, orig in title_norm_to_orig.items():
        ratio = SequenceMatcher(None, norm_text, norm_title).ratio()
        if ratio > best[1]:
            best = (orig, ratio)
    return best


def _build_map_from_docx(docx_path: Path, titles: List[str]) -> Dict[str, str]:
    title_norm_to_orig = {_normalize_title(t): t for t in titles}
    current_key: str | None = None
    buffer: List[str] = []
    result: Dict[str, str] = {}

    try:
        doc = Document(str(docx_path))
    except Exception as e:
        logging.exception("Failed to open hypotheses DOCX: %s", e)
        return {}

    paragraphs = [p.text.strip() for p in doc.paragraphs if (p.text or "").strip()]

    for p in paragraphs:
        norm = _normalize_title(p)
        matched_title, score = _best_title_match(norm, title_norm_to_orig)
        if matched_title and score >= 0.85:
            # flush previous
            if current_key is not None:
                result[current_key] = "\n".join(buffer).strip()
            current_key = matched_title
            buffer = []
        else:
            if current_key is not None:
                buffer.append(p)

    if current_key is not None:
        result[current_key] = "\n".join(buffer).strip()

    return result


def _should_rebuild(docx_path: Path, cache_path: Path) -> bool:
    if not cache_path.exists():
        return True
    try:
        return os.path.getmtime(docx_path) > os.path.getmtime(cache_path)
    except Exception:
        return True


def ensure_hypotheses_map() -> Dict[str, str]:
    titles = load_hypotheses_list()
    if not titles:
        return {}
    if HYPOTHESES_DOCX.exists() and _should_rebuild(HYPOTHESES_DOCX, HYPOTHESES_MAP_JSON):
        mapping = _build_map_from_docx(HYPOTHESES_DOCX, titles)
        try:
            HYPOTHESES_MAP_JSON.write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8")
        except Exception:
            logging.exception("Failed to write hypotheses_map.json")
        return mapping
    # load cache
    try:
        if HYPOTHESES_MAP_JSON.exists():
            return json.loads(HYPOTHESES_MAP_JSON.read_text(encoding="utf-8"))
    except Exception:
        logging.exception("Failed to read hypotheses_map.json")
    # fallback: rebuild if possible
    if HYPOTHESES_DOCX.exists():
        return _build_map_from_docx(HYPOTHESES_DOCX, titles)
    return {}


def get_description_for_title(title: str) -> str | None:
    mapping = ensure_hypotheses_map()
    if not mapping:
        return None
    # direct match
    if title in mapping:
        return mapping[title]
    # fuzzy: normalize
    tnorm = _normalize_title(title)
    # best ratio over existing keys
    try:
        from difflib import SequenceMatcher
        best_k = None
        best_ratio = 0.0
        for k in mapping.keys():
            ratio = SequenceMatcher(None, tnorm, _normalize_title(k)).ratio()
            if ratio > best_ratio:
                best_ratio = ratio
                best_k = k
        if best_k and best_ratio >= 0.75:
            return mapping[best_k]
    except Exception:
        pass
    return None


def _format_description_html(text: str) -> str:
    # Normalize paragraphs, highlight sub-headers, convert bullets to readable dots
    lines = [ln.strip() for ln in text.splitlines()]
    out: List[str] = []
    for raw in lines:
        if not raw:
            out.append("")
            continue
        # bullets
        if raw.startswith(("- ", "• ", "— ", "· ")):
            content = raw[2:].strip()
            out.append(f"• {_html.escape(content)}")
            continue
        # numeric bullets
        import re as _re
        if _re.match(r"^\d+[\.)]\s+", raw):
            content = _re.sub(r"^\d+[\.)]\s+", "", raw).strip()
            out.append(f"• {_html.escape(content)}")
            continue
        # section header ending with ':'
        if raw.endswith(":") and len(raw) <= 120:
            hdr = raw[:-1].strip()
            out.append(f"<b>{_html.escape(hdr)}</b>:")
            continue
        out.append(_html.escape(raw))
    # collapse extra blank lines
    result_lines: List[str] = []
    prev_blank = False
    for ln in out:
        blank = ln == ""
        if blank and prev_blank:
            continue
        result_lines.append(ln)
        prev_blank = blank
    return "\n".join(result_lines).strip()


def get_description_formatted_html(title: str) -> str | None:
    """Return description converted to Telegram-safe HTML with bullets and headers."""
    d = get_description_for_title(title)
    if not d:
        return None
    return _format_description_html(d)


# ---- Full card extraction (scenario block containing hypothesis) ----

def _read_doc_paragraphs(docx_path: Path) -> List[str]:
    try:
        doc = Document(str(docx_path))
        return [p.text.strip() for p in doc.paragraphs if (p.text or "").strip()]
    except Exception:
        return []


def _is_header_line(text: str) -> bool:
    if len(text) < 5:
        return False
    t = text.replace("—", "-")
    # Header if mostly uppercase and contains no lowercase letters
    import re as _re
    if _re.match(r"^[A-ZА-ЯЁ0-9 ,()\-–—]+$", t):
        return True
    return False


def _split_into_blocks(pars: List[str]) -> List[List[str]]:
    blocks: List[List[str]] = []
    current: List[str] = []
    for line in pars:
        if _is_header_line(line):
            if current:
                blocks.append(current)
            current = [line]
        else:
            if not current:
                current = [line]
            else:
                current.append(line)
    if current:
        blocks.append(current)
    return blocks


def _block_contains_hypothesis(block: List[str], title: str) -> bool:
    norm_title = _normalize_title(title)
    txt = "\n".join(block)
    if norm_title in _normalize_title(txt):
        return True
    # Heuristic: scan lines after a "ГИПОТЕЗ" header
    import re as _re
    in_hyp = False
    for ln in block:
        if not in_hyp and _re.search(r"^\s*ГИПОТЕЗ", ln.upper()):
            in_hyp = True
            continue
        if in_hyp and ln.endswith(":"):
            # next section begins
            break
        if in_hyp:
            if norm_title in _normalize_title(ln):
                return True
            try:
                from difflib import SequenceMatcher
                if SequenceMatcher(None, _normalize_title(ln), norm_title).ratio() >= 0.8:
                    return True
            except Exception:
                pass
    return False


def get_full_card_text_for_hypothesis(title: str) -> str | None:
    if not HYPOTHESES_DOCX.exists():
        return get_description_for_title(title)
    pars = _read_doc_paragraphs(HYPOTHESES_DOCX)
    if not pars:
        return get_description_for_title(title)
    blocks = _split_into_blocks(pars)
    for blk in blocks:
        try:
            if _block_contains_hypothesis(blk, title):
                return "\n".join(blk).strip()
        except Exception:
            continue
    return get_description_for_title(title)


def get_full_card_formatted_html(title: str) -> str | None:
    txt = get_full_card_text_for_hypothesis(title)
    if not txt:
        return None
    return _format_description_html(txt)


# ---- JSON cache of full blocks and cosine matching ----

def _should_rebuild_full(docx_path: Path, cache_path: Path) -> bool:
    if not cache_path.exists():
        return True
    try:
        return os.path.getmtime(docx_path) > os.path.getmtime(cache_path)
    except Exception:
        return True


def _blocks_to_json(blocks: List[List[str]]) -> List[Dict[str, str]]:
    items: List[Dict[str, str]] = []
    for blk in blocks:
        if not blk:
            continue
        title = blk[0].strip()
        text = "\n".join(blk[1:]).strip()
        items.append({"title": title, "text": text})
    return items


def ensure_full_cards_json() -> List[Dict[str, str]]:
    # restore persistent cache for reliability and debugging
    if HYPOTHESES_DOCX.exists() and _should_rebuild_full(HYPOTHESES_DOCX, HYPOTHESES_FULL_JSON):
        pars = _read_doc_paragraphs(HYPOTHESES_DOCX)
        blocks = _split_into_blocks(pars)
        data = _blocks_to_json(blocks)
        try:
            HYPOTHESES_FULL_JSON.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        except Exception:
            logging.exception("Failed to write hypotheses_full.json")
        return data
    try:
        if HYPOTHESES_FULL_JSON.exists():
            return json.loads(HYPOTHESES_FULL_JSON.read_text(encoding="utf-8"))
    except Exception:
        logging.exception("Failed to read hypotheses_full.json")
    if HYPOTHESES_DOCX.exists():
        pars = _read_doc_paragraphs(HYPOTHESES_DOCX)
        blocks = _split_into_blocks(pars)
        return _blocks_to_json(blocks)
    return []


def _tokenize(text: str) -> List[str]:
    import re as _re
    t = text.lower().replace("ё", "е").replace("—", "-").replace("–", "-")
    return _re.findall(r"[a-zа-я0-9][a-zа-я0-9\-]{1,}", t)


def _idf(corpus_tokens: List[List[str]]) -> Dict[str, float]:
    import math
    df: Dict[str, int] = {}
    N = len(corpus_tokens)
    for doc in corpus_tokens:
        for tok in set(doc):
            df[tok] = df.get(tok, 0) + 1
    return {tok: math.log((1 + N) / (1 + c)) + 1.0 for tok, c in df.items()}


def _tf(tokens: List[str]) -> Dict[str, float]:
    tf: Dict[str, float] = {}
    for t in tokens:
        tf[t] = tf.get(t, 0.0) + 1.0
    return tf


def _cosine(a: Dict[str, float], b: Dict[str, float]) -> float:
    import math
    # dot
    dot = 0.0
    for k, v in a.items():
        if k in b:
            dot += v * b[k]
    na = math.sqrt(sum(v * v for v in a.values()))
    nb = math.sqrt(sum(v * v for v in b.values()))
    if na == 0.0 or nb == 0.0:
        return 0.0
    return dot / (na * nb)


def get_best_full_card_by_cosine(query: str) -> Dict[str, str] | None:
    data = ensure_full_cards_json()
    if not data:
        return None
    docs_tokens: List[List[str]] = []
    for item in data:
        # weight title tokens more by duplicating
        toks = _tokenize(item.get("title", "")) * 3 + _tokenize(item.get("text", ""))
        docs_tokens.append(toks)
    idf = _idf(docs_tokens)
    # build doc vectors
    docs_vec: List[Dict[str, float]] = []
    for toks in docs_tokens:
        tf = _tf(toks)
        vec = {t: tf[t] * idf.get(t, 0.0) for t in tf}
        docs_vec.append(vec)
    # query vector
    q_tokens = _tokenize(query)
    q_tf = _tf(q_tokens)
    q_vec = {t: q_tf[t] * idf.get(t, 0.0) for t in q_tf}
    # argmax cosine
    best_i = -1
    best_s = -1.0
    for i, vec in enumerate(docs_vec):
        s = _cosine(q_vec, vec)
        if s > best_s:
            best_s = s
            best_i = i
    return data[best_i] if best_i >= 0 else None


def format_card_fields_to_html(title: str, text: str) -> str:
    """Public wrapper to format a card given separate title and body text."""
    combined = (title or "").strip()
    body = (text or "").strip()
    if body:
        combined = (combined + "\n" + body) if combined else body
    return _format_description_html(combined)


# ---- Exact map: hypothesis title -> full block text (from DOCX) ----

def _build_cards_index_by_titles() -> Dict[str, str]:
    """Strict slicing: for each title from hypotheses.txt, take text from its first
    occurrence in the DOCX up to the first occurrence of the NEXT title.
    Returns a map: exact title -> full block (title + body)."""
    titles = load_hypotheses_list()
    if not titles:
        return {}
    pars = _read_doc_paragraphs(HYPOTHESES_DOCX) if HYPOTHESES_DOCX.exists() else []
    if not pars:
        return {}
    norm_title_list = [_normalize_title(t) for t in titles]
    # also store left-side (before dash) variants to match concise headers
    norm_left_list = [_normalize_title(t.split("—")[0].split("-")[0]) for t in titles]
    # find start index for each title by first paragraph containing it
    starts: List[int] = []
    for idx_t, nt in enumerate(norm_title_list):
        idx = -1
        for i, p in enumerate(pars):
            pn = _normalize_title(p)
            nl = norm_left_list[idx_t]
            if (nt and nt in pn) or (nl and nl in pn):
                idx = i
                break
        starts.append(idx)
    # Build slices
    result: Dict[str, str] = {}
    for i, t in enumerate(titles):
        s = starts[i]
        if s == -1:
            # if title not found, skip
            continue
        # find next start > s
        e = len(pars)
        for j in range(i + 1, len(starts)):
            if starts[j] != -1 and starts[j] > s:
                e = starts[j]
                break
        block_lines = pars[s:e]
        block_text = "\n".join(block_lines).strip()
        # ensure block begins with the exact title line if present; else prepend title
        if block_lines and _normalize_title(block_lines[0]) != _normalize_title(t):
            block_text = f"{t}\n{block_text}"
        result[t] = block_text
    return result


def ensure_cards_index_json() -> Dict[str, str]:
    titles = load_hypotheses_list()
    need_rebuild = (
        (not HYPOTHESES_CARDS_JSON.exists())
        or (HYPOTHESES_DOCX.exists() and _should_rebuild_full(HYPOTHESES_DOCX, HYPOTHESES_CARDS_JSON))
    )
    if not need_rebuild:
        try:
            existing = json.loads(HYPOTHESES_CARDS_JSON.read_text(encoding="utf-8")) if HYPOTHESES_CARDS_JSON.exists() else {}
        except Exception:
            existing = {}
        # if existing doesn't have all 23 keys -> rebuild
        if not isinstance(existing, dict) or len(existing.keys()) != len(titles):
            need_rebuild = True
        else:
            return existing

    # Build fresh index strictly by titles order
    base = _build_cards_index_by_titles()
    # Fill missing using cosine or block-finder fallbacks to guarantee all 23 keys
    filled: Dict[str, str] = {}
    for t in titles:
        txt = base.get(t)
        if not txt:
            # try cosine
            bc = get_best_full_card_by_cosine(t)
            if bc:
                txt = (bc.get("title", "") + "\n" + bc.get("text", "")).strip()
        if not txt:
            # try section that contains title
            txt = get_full_card_text_for_hypothesis(t) or ""
        # ensure prepend title if not present
        if txt and _normalize_title(txt.split("\n", 1)[0]) != _normalize_title(t):
            txt = f"{t}\n{txt}"
        filled[t] = txt

    try:
        HYPOTHESES_CARDS_JSON.write_text(json.dumps(filled, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        logging.exception("Failed to write hypotheses_cards.json")
    return filled


def get_card_by_hypothesis_title_html(title: str) -> str | None:
    idx = ensure_cards_index_json()
    if not idx:
        return None
    # exact first
    if title in idx:
        return _format_description_html(idx[title])
    # normalized lookup
    tn = _normalize_title(title)
    for k, v in idx.items():
        if _normalize_title(k) == tn:
            return _format_description_html(v)
    # fuzzy over 23 titles
    try:
        from difflib import SequenceMatcher
        best_k = None
        best_ratio = 0.0
        for k in idx.keys():
            r = SequenceMatcher(None, tn, _normalize_title(k)).ratio()
            if r > best_ratio:
                best_ratio = r
                best_k = k
        if best_k and best_ratio >= 0.65:
            return _format_description_html(idx[best_k])
    except Exception:
        pass
    return None


def _extract_section(text: str, header: str) -> str:
    lines = [ln.rstrip() for ln in text.splitlines()]
    out: list[str] = []
    capture = False
    hdr_upper = header.upper().strip().rstrip(":")
    for ln in lines:
        ln_stripped = ln.strip()
        if not capture:
            if ln_stripped.upper().startswith(hdr_upper + ":"):
                # start capturing from the same line (after colon)
                after = ln_stripped.split(":", 1)[1].strip() if ":" in ln_stripped else ""
                if after:
                    out.append(after)
                capture = True
            continue
        # stop on next all-caps section label ending with ':'
        if ln_stripped and ln_stripped == ln_stripped.upper() and ln_stripped.endswith(":"):
            break
        out.append(ln)
    return "\n".join([ln for ln in out]).strip()


def get_card_description_only_html(title: str) -> str | None:
    idx = ensure_cards_index_json()
    block = None
    if title in idx:
        block = idx[title]
    else:
        tn = _normalize_title(title)
        for k, v in idx.items():
            if _normalize_title(k) == tn:
                block = v
                break
    if not block:
        return None
    desc = _extract_section(block, "ОПИСАНИЕ")
    if not desc:
        return None
    return _format_description_html(desc)


